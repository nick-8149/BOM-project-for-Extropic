import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any

def create_excel_report(
    title: str,
    headers: List[str],
    rows: List[List[Any]],
    summary_data: Dict[str, Any] = None
) -> bytes:
    """
    Generates a professionally-styled Excel spreadsheet from raw grid rows.
    Uses openpyxl to apply harmonize premium dark colors, custom typography,
    and automatic column width sizing.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    
    # Enable grid lines visibility
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions (Harmonious engineering palette)
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
    section_font = Font(name="Segoe UI", size=11, bold=True, color="334155")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="000000")
    bold_data_font = Font(name="Segoe UI", size=10, bold=True, color="000000")
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    summary_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    double_bottom_border = Border(bottom=Side(border_style="double", color="334155"), top=thin_border_side)

    # 1. Report Title Block
    ws.cell(row=1, column=1, value=title).font = title_font
    ws.row_dimensions[1].height = 30
    
    current_row = 3

    # 2. Add Summary KPI Block if provided
    if summary_data:
        ws.cell(row=current_row, column=1, value="Report Summary Metrics").font = section_font
        ws.row_dimensions[current_row].height = 20
        current_row += 1
        
        for k, v in summary_data.items():
            cell_k = ws.cell(row=current_row, column=1, value=k)
            cell_v = ws.cell(row=current_row, column=2, value=v)
            
            cell_k.font = bold_data_font
            cell_k.fill = summary_fill
            cell_k.border = thin_border
            
            cell_v.font = data_font
            cell_v.fill = summary_fill
            cell_v.border = thin_border
            current_row += 1
            
        current_row += 2 # Spacer rows

    # 3. Add Table Headers
    ws.row_dimensions[current_row].height = 25
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    current_row += 1

    # 4. Add Data Rows
    for r_idx, row in enumerate(rows):
        ws.row_dimensions[current_row].height = 20
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            
            # Align numeric columns
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    # 5. Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            # Skip title row length calculations
            if cell.row == 1:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save workbook to memory bytes
    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    return file_stream.getvalue()
